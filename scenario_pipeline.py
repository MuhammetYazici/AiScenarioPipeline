import glob
import json
import os
import argparse
import sys
import re
import subprocess
import time
import docx
import openpyxl
from dotenv import load_dotenv
from google import genai
from playwright.sync_api import sync_playwright

# .env dosyasındaki gizli anahtarı okur (GEMINI_API_KEY)
load_dotenv()

# Google API Client
client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

# Java projesinin ana kök dizini (pom.xml dosyasının bulunduğu yer)
JAVA_PROJECT_BASE = r"C:\Users\Muhammet  YAZICI\IdeaProjects\Proje250824_1\aiStagerPlaywright"
# Java test kaynak kodlarının kök dizini
JAVA_PROJECT_ROOT = os.path.join(JAVA_PROJECT_BASE, r"src\test\java")
FEATURE_DIR = "aiStagerPlaywright/src/test/java/FeatureFiles"
FEATURE_FILE_PATH = os.path.join(FEATURE_DIR, "generated_scenario.feature")
CONFIG_PROPERTIES_PATH = "aiStagerPlaywright/src/test/resources/config.properties"


# --- 0. GÜVENLİ API İSTEK KATMANI (RETRY / 503 DAYANIKLILIĞI) ---
def safe_generate_content(prompt: str, max_retries: int = 5) -> str:
    """
    Google API tarafında anlık 503 UNAVAILABLE veya 429 RATE LIMIT hatası olduğunda
    akışın çökmesini engeller; bekleme süresini katlayarak isteği tekrar dener.
    """
    models_to_try = ["gemini-flash-latest",
                     "gemini-3.5-flash"
                     ]

    for model_name in models_to_try:
        for attempt in range(1, max_retries + 1):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                )
                return response.text
            except Exception as e:
                err_msg = str(e)
                if any(err_key in err_msg for err_key in ["503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED"]):
                    wait_time = attempt * 5
                    print(f">> API Sunucu Yoğunluğu ({model_name}): {err_msg[:200]}")
                    print(f">> {wait_time} saniye bekleniyor (Deneme {attempt}/{max_retries})...")
                    time.sleep(wait_time)
                else:
                    print(f">> {model_name} hata verdi ({err_msg[:200]}...), alternatif modele geçiliyor...")
                    break

    raise RuntimeError("Google API sunucuları şu anda yanıt veremedi. Lütfen birkaç dakika sonra tekrar deneyin.")



# --- 1. DOKÜMAN OKUYUCULAR VE USER STORY ÇIKARMA ---
def extract_text_from_docx(file_path: str) -> str:
    """Word (.docx) dosyasındaki tüm paragrafları düz metin olarak okur."""
    doc = docx.Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])


def extract_text_from_excel(file_path: str) -> str:
    """Excel (.xlsx) dosyasındaki tüm sayfaları ve dolu satırları metne çevirir."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    full_text = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        full_text.append(f"\n--- Sayfa: {sheet_name} ---")
        for row in sheet.iter_rows(values_only=True):
            row_values = [
                str(cell).strip()
                for cell in row
                if cell is not None and str(cell).strip() != ""
            ]
            if row_values:
                full_text.append(" | ".join(row_values))
    return "\n".join(full_text)


def extract_content(file_path: str) -> str:
    """Dosya uzantısına göre doğru okuyucu fonksiyonu çalıştırır."""
    if file_path.endswith(".docx"):
        return extract_text_from_docx(file_path)
    elif file_path.endswith(".xlsx") or file_path.endswith(".xls"):
        return extract_text_from_excel(file_path)
    return ""


class PlaywrightPlanner:
    def __init__(self, base_url: str = "https://aistager.ai/tr"):
        self.base_url = base_url

    def resolve_target_url(self, feature_text: str) -> str:
        """Senaryodaki adımları inceleyerek açılması gereken hedef URL'yi belirler."""
        lower_text = feature_text.lower()
        if "kayıt" in lower_text or "register" in lower_text:
            return f"{self.base_url}/register"
        elif "giriş" in lower_text or "login" in lower_text:
            return f"{self.base_url}/login"
        elif any(k in lower_text for k in ["görsel", "staging", "sanal sahneleme"]):
            return f"{self.base_url}/virtual-staging"
        elif "api" in lower_text:
            return f"{self.base_url}/api"
        return self.base_url

    def discover_interactive_elements(self, target_url: str) -> list[dict]:
        """
        Sayfayı headless açar; erişilebilirlik (AOM) ve semantik ağaçtaki
        interaktif elementleri (button, textbox, link, select, checkbox) ayıklar.
        """
        print(f"\n>> 🎭 [Planner Agent] Canlı sayfa taranıyor: {target_url}")
        elements = []

        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                context = browser.new_context(
                    viewport={"width": 1600, "height": 900},
                    user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
                )
                page = context.new_page()

                try:
                    page.goto(target_url, wait_until="domcontentloaded", timeout=15000)
                    page.wait_for_timeout(2000)
                except Exception as nav_err:
                    print(f">> [Planner Uyarı]: Sayfa geç yüklendi: {nav_err}")

                # DOM üzerinden temiz AOM haritası çekilir
                elements = page.evaluate("""
                                         () => {
                                             const nodes = document.querySelectorAll(
                                                 "button, input, select, textarea, a, [role='button'], [role='textbox'], [role='checkbox'], [role='link']"
                                             );
                                             return Array.from(nodes).map(el => {
                                                 const ariaLabel = el.getAttribute('aria-label') || '';
                                                 const role = el.getAttribute('role') || el.tagName.toLowerCase();
                                                 const placeholder = el.getAttribute('placeholder') || '';
                                                 const name = el.getAttribute('name') || '';
                                                 const type = el.getAttribute('type') || '';
                                                 const text = (el.innerText || el.value || '').trim().replace(/\\s+/g, ' ').substring(0, 50);

                                                 return {
                                                     tag: el.tagName.toLowerCase(),
                                                     role: role,
                                                     type: type,
                                                     accessibleName: ariaLabel || placeholder || text || name,
                                                     placeholder: placeholder,
                                                     text: text
                                                 };
                                             }).filter(n => n.accessibleName || n.text);
                                         }
                                         """)
                browser.close()
                print(f">> 🎭 [Planner Agent] {len(elements)} adet interaktif element yakalandı.")
        except Exception as e:
            print(f">> [Planner Hatası]: Sayfa keşfi yapılamadı: {e}")

        return elements

    def create_execution_plan(self, feature_text: str, scanned_elements: list[dict], safe_generate_func) -> str:
        """
        Keşfedilen elementler ile Feature adımlarını eşleştirerek
        Generator'ın kullanacağı kesin Eylem Planını (Action Plan) çıkarır.
        """
        print(">> 📋 [Planner Agent] Senaryo için Eylem Planı (Execution Plan) oluşturuluyor...")

        prompt = f"""
            Sen kıdemli bir Test Otomasyon Mimarı ve Playwright Uzmanısın.
            Aşağıda test edilecek Gherkin senaryosu ve canlı sayfadan Playwright ile çıkarılan interaktif elementler listesi var.

            GÖREVİN:
            Senaryodaki adımları canlı elementlerle eşleştirerek Generator'ın (Kod Üretici) doğrudan koda çevirebileceği bir Eylem Planı (Action Plan) oluştur.

            KURALLAR:
            1. Her adım için kullanılacak Playwright seçicisini belirle:
                - Butonlar için: getByRole(AriaRole.BUTTON, ...)
                - Metin kutuları için: getByRole(AriaRole.TEXTBOX, ...) veya getByPlaceholder(...)
                - Linkler için: getByRole(AriaRole.LINK, ...)
            2. Sayfada olmayan elementler için varsayımda bulunma, en yakın semantik eşleşmeyi seç.
            3. Çıktıyı temiz bir JSON formatında ver.

        Gherkin Senaryosu:
        {feature_text}

        Canlı Elementler (İlk 60 element):
        {json.dumps(scanned_elements[:60], ensure_ascii=False, indent=2)}

        Çıktı Formatı (Sadece bu JSON bloğunu ver):
        ```json
        [
        {{
            "step": "Gherkin adım metni",
            "action": "click | fill | verify | select | check",
            "locator_strategy": "getByRole | getByPlaceholder | getByText",
            "role": "BUTTON | TEXTBOX | LINK | vb.",
            "name_or_text": "Element adı veya placeholder",
            "value": "Varsa girilecek veri veya null"
        }}
        ]
    """
        raw_plan = safe_generate_func(prompt)
        clean_plan = re.sub(r"^```[a-zA-Z]*\n|```$", "", raw_plan.strip()).strip()
        return clean_plan

def convert_doc_to_user_story(raw_text: str) -> str:
    """Ham doküman metnini analiz edip kurallı bir User Story çıkarır."""
    prompt = f"""
    Sen kıdemli bir Business Analyst ve QA Mimarı rolündesin.
    Aşağıdaki ham doküman metnini/tablosunu incele; sohbet kalıntılarını ve gereksiz detayları ayıkla.
    Doğrudan test edilebilir iş kurallarını (Business Rules) ve senaryo hedeflerini içeren net bir kullanıcı hikayesi (User Story) çıkar.

    Metin:
    {raw_text}
    """
    return safe_generate_content(prompt)


def generate_bdd_feature(user_story: str) -> str:
    """
    User story ve iş kurallarından doğrudan kurallara tam uyan Gherkin Feature üretir.
    """
    prompt = f"""
Sen kıdemli bir Test Mimarı ve BDD Uzmanısın.
Aşağıdaki kullanıcı hikayesi ve iş kurallarını incele. Tüm pozitif, negatif ve sınır (edge case) 
durumları kapsayan eksiksiz bir Cucumber '.feature' metni üret.

KESİN SÖZDİZİMİ KURALLARI:
1. Dosyada SADECE BİR (1) adet 'Feature:' olmalıdır.
2. SADECE BİR (1) adet ortak 'Background:' olabilir. Alt sayfalara gitme adımlarını (login, register vb.) Background'a değil ilgili Scenario'nun ilk satırına yaz.
3. Gherkin anahtar kelimeleri İNGİLİZCE (Feature, Background, Scenario, Given, When, Then, And), adım metinleri TÜRKÇE olmalıdır.
4. Adım metinlerinde ASLA parantez () KULLANMA. Çift tırnağı ("") SADECE dinamik veriler ve URL'ler için kullan.
5. Markdown formatı (```gherkin) KULLANMA, doğrudan saf metin ver.

Kullanıcı Hikayesi ve İş Kuralları:
{user_story}
"""
    raw_response = safe_generate_content(prompt)
    clean_text = raw_response.strip()
    clean_text = re.sub(r"^```[a-zA-Z]*\n|```$", "", clean_text).strip()
    clean_text = re.sub(r"^\s*#\s*language:\s*[a-zA-Z_-]+\s*\n", "", clean_text, flags=re.MULTILINE)
    return clean_text


def clean_gherkin_syntax(raw_feature: str) -> str:
    """
    LLM'in kaçırabileceği parantezleri, hatalı tırnakları ve URL kaçışlarını
    sıfır token harcayarak Python regex ile deterministik temizler.
    """
    def normalize_step_line(match):
        keyword = match.group(1)
        body = match.group(2).strip()

        # Parantezli ifadeleri ve tırnak karmaşalarını düzelt
        body = re.sub(r'\(\s*("[^"]*")\s*\)', r'\1', body)
        body = re.sub(r'\s*\([^)]*\)', '', body)
        body = re.sub(r'\s{2,}', ' ', body).strip()

        # URL geçen yerleri tırnak içine al
        body = re.sub(r'(?<!")([a-zA-Z0-9.-]+\.[a-zA-Z]{2,}(?:/[^\s"]*)?)(?!")', r'"\1"', body)
        return f"{keyword} {body}"

    cleaned = re.sub(r'^(Given|When|Then|And|But)\s+(.*)$', normalize_step_line, raw_feature, flags=re.MULTILINE)
    return cleaned.strip()



# --- 2.5 CUCUMBER CLI RESMİ SÖZLEŞME (CONTRACT) MOTORU ---
def extract_official_cucumber_snippets(feature_filename: str) -> str:
    """
    Cucumber dry-run çalıştırarak ilgili feature'a ait resmi, hatasız
    Java metot imzalarını (Snippets) çeker. Bu liste LLM için bağlayıcı bir sözleşmedir.
    """
    print(f"\n>> 📐 [Cucumber CLI] {feature_filename} için resmi Java sözleşmesi çıkarılıyor...")

    feature_rel_path = f"src/test/java/FeatureFiles/{feature_filename}"
    cmd = f'mvn test -Dcucumber.features="{feature_rel_path}" -Dcucumber.execution.dry-run=true -Dcucumber.plugin="summary"'

    res = subprocess.run(
        cmd,
        cwd=JAVA_PROJECT_BASE,
        shell=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="ignore",
    )

    output = res.stdout
    marker = "You can implement missing steps with the snippets below:"

    if marker in output:
        raw_snippets = output.split(marker)[1]
        if "```" in raw_snippets:
            raw_snippets = raw_snippets.split("```")[0]
        clean_snippets = re.sub(r"\[INFO\]|\[WARNING\]", "", raw_snippets).strip()
        print(">> 📐 [Cucumber CLI] Başarılı: Resmi Java şablonları doğrudan Cucumber motorundan alındı.")
        return clean_snippets
    else:
        print(">> [Cucumber CLI] Tüm adımlar zaten tanımlı veya şablon doğrudan gerekmedi.")
        return ""



# --- 4. DİNAMİK JAVA KODU ÜRETİCİSİ (LEAD SDET + SÖZLEŞME ENTEGRELİ) ---
class JavaCodeGenerator:
    def __init__(self, java_project_root: str, safe_generate_func):
        self.java_project_root = java_project_root
        self.safe_generate = safe_generate_func

    def generate_and_save_runner(self, class_name_hint: str):
        """TestNG Runner sınıfını deterministik şablonla üretir (Sıfır LLM / Sıfır Hata)."""
        runner_dir = os.path.join(self.java_project_root, "Runners")
        os.makedirs(runner_dir, exist_ok=True)
        runner_file = os.path.join(runner_dir, f"{class_name_hint}Runner.java")

        runner_code = f"""package Runners;

import io.cucumber.testng.AbstractTestNGCucumberTests;
import io.cucumber.testng.CucumberOptions;

@CucumberOptions(
    features = "src/test/java/FeatureFiles",
    glue = {{"StepDefinations", "Hooks"}},
    plugin = {{
        "pretty",
        "html:target/cucumber-reports.html",
        "json:target/cucumber.json"
    }}
)
public class {class_name_hint}Runner extends AbstractTestNGCucumberTests {{
}}
"""
        with open(runner_file, "w", encoding="utf-8") as f:
            f.write(runner_code.strip())
        print(f">> 🚀 [Generator] Deterministik Runner oluşturuldu: {runner_file}")

    def generate_page_class(self, class_name_hint: str, execution_plan: str) -> str:
        """Planner'ın eylem planını alıp BasePage'i extends eden Page sınıfını üretir."""
        print(f">> 📄 [Generator] {class_name_hint}Page.java sınıfı üretiliyor...")

        prompt = f"""
Sen kıdemli bir Playwright Java mühendisisin.
Aşağıdaki Eylem Planını (Execution Plan) kullanarak Pages/{class_name_hint}Page.java sınıfını yaz.

MİMARİ KURALLAR:
1. Sınıf MUTLAKA 'public class {class_name_hint}Page extends BasePage' olmalıdır.
2. 'package Pages;' altında olmalıdır.
3. Parametresiz default constructor koy (super(); çağrısı zaten BasePage içindedir, 'this.page = PD.getPage();' yazma!).
4. BASEPAGE METOT İMZALARI (ÇOK KRİTİK):
   - 'myClick(Locator element)' -> Sadece Locator alır, ASLA String almaz! (Örn: myClick(getLoginButton());)
   - 'mySendKeys(Locator element, String text)' -> İlk parametre Locator, ikinci String!
   - 'myCheckBox(Locator element)' -> Sadece Locator alır, ikinci bir boolean parametre KESİNLİKLE VERME!
   - 'verifyContainsText(Locator element, String text)' -> İlk parametre Locator!
5. ASLA Driver veya ReusableMethod sınıfı çağırma!
6. Playwright seçicilerini AOM standartlarında yaz:
   - page.getByRole(AriaRole.BUTTON, new Page.GetByRoleOptions().setName("..."))
   - page.getByRole(AriaRole.TEXTBOX, new Page.GetByRoleOptions().setName("..."))
   - page.getByPlaceholder("...")
7. Gerekli importları ekle:
   - import Pages.BasePage;
   - import com.microsoft.playwright.options.AriaRole;
   - import com.microsoft.playwright.Page;
   - import com.microsoft.playwright.Locator;

Eylem Planı (Execution Plan):
{execution_plan}

ÇIKTI KURALI:
Sadece saf Java kodunu ver. Selamlama veya markdown formatı (```java) KULLANMA.
"""
        raw_code = self.safe_generate(prompt)
        clean_code = re.sub(r"^```[a-zA-Z]*\n|```$", "", raw_code.strip()).strip()

        page_dir = os.path.join(self.java_project_root, "Pages")
        os.makedirs(page_dir, exist_ok=True)
        page_file = os.path.join(page_dir, f"{class_name_hint}Page.java")
        with open(page_file, "w", encoding="utf-8") as f:
            f.write(clean_code)
        print(f">> 📄 [Generator] Page sınıfı diske yazıldı: {page_file}")
        return clean_code

    def generate_step_class(self, class_name_hint: str, generated_page_code: str, official_snippets: str,
                            feature_text: str) -> str:
        """Cucumber sözleşmesini, üretilen Page sınıfını ve TestContext'i birleştirerek Step sınıfını üretir."""
        print(f">> 👣 [Generator] {class_name_hint}Steps.java sınıfı üretiliyor...")

        prompt = f"""
Sen kıdemli bir Test Otomasyon Mühendisisin.
Aşağıdaki bileşenleri kullanarak StepDefinations/{class_name_hint}Steps.java sınıfını oluştur.

KESİN SINIRLAR VE YASAKLAR:
1. JUnit KESİNLİKLE YASAKTIR! Asla 'org.junit' import etme. Doğrulamalar için 'org.testng.Assert' kullan.
2. ASLA 'DriverFactory', 'Driver' veya 'ReusableMethods' çağırma!
3. 'context.getPage()' METODU YOKTUR! Sayfaya erişmek için 'Utilities.PD.getPage()' kullan.
4. Adımları yürütürken doğrudan ilgili Page sınıfının metotlarını çağır: 'page.metotAdi()'.

MİMARİ VE PICO-CONTAINER KURALLARI (ÇOK KRİTİK):
1. 'package StepDefinations;' altında olmalıdır.
2. Sınıf başında nesneleri tanımla:
   - private final {class_name_hint}Page page;
   - private final TestContext context;
3. CONSTRUCTOR INJECTION KURALI:
   PicoContainer için constructor parametre olarak TestContext almalıdır:
   public {class_name_hint}Steps(TestContext context) {{
       this.context = context;
       this.page = new {class_name_hint}Page();
   }}
4. Asla 'new TestContext()' veya static değişken kullanma! Senaryo verilerini context.set("key", val) ve context.get("key", String.class) ile paylaş.
5. CUCUMBER RESMİ SÖZLEŞMESİNE KESİNLİKLE SADIK KAL:
   Resmi imzaların metinlerine (@Given/@When/@Then) ve parametrelerine dokunma, sadece süslü parantezlerin {{ }} içini doldur.
   İçeride ilgili '{class_name_hint}Page' metotlarını çağır.
6. Gövdesi boş veya sadece yorum olan metot kesinlikle yasaktır!
7. ADIM ANOTASYONLARI VE PARAMETRE KURALI (ÇOK KRİTİK):
   - Feature dosyasında çift tırnak ("...") içinde geçen tüm URL, metin veya dinamik değerleri Step metot anotasyonunda MUTLAKA '{{string}}' olarak tanımla!
   - ASLA anotasyon içine kaçışlı tırnak (\") yazma.
   - ÖRNEK (YANLIŞ): @Then("Kullanıcı doğrudan \"/tr/ai-virtual-tour\" adresine yönlendirilmelidir")
   - ÖRNEK (DOĞRU):   @Then("Kullanıcı doğrudan {{string}} adresine yönlendirilmelidir")
     public void kullaniciDogrudanAdresineYonlendirilmelidir(String url) { ... }

CUCUMBER RESMİ İMZALARI (SÖZLEŞME):
{official_snippets if official_snippets else "Resmi sözleşme boşsa Feature adımlarını baz al."}

KULLANILACAK PAGE SINIFI KODU:
{generated_page_code}

FEATURE METNİ:
{feature_text}

ÇIKTI KURALI:
Sadece saf Java kodunu ver. Markdown formatı (```java) veya açıklama ekleme.
"""
        raw_code = self.safe_generate(prompt)
        clean_code = re.sub(r"^```[a-zA-Z]*\n|```$", "", raw_code.strip()).strip()

        step_dir = os.path.join(self.java_project_root, "StepDefinations")
        os.makedirs(step_dir, exist_ok=True)
        step_file = os.path.join(step_dir, f"{class_name_hint}Steps.java")
        with open(step_file, "w", encoding="utf-8") as f:
            f.write(clean_code)
        print(f">> 👣 [Generator] Step sınıfı diske yazıldı: {step_file}")
        return clean_code


class SelfHealer:
    def __init__(self, java_project_base: str, java_project_root: str, safe_generate_func):
        self.java_project_base = java_project_base
        self.java_project_root = java_project_root
        self.safe_generate = safe_generate_func

    def check_and_heal_missing_cucumber_steps(self, class_name_hint: str) -> bool:
        """Cucumber dry-run yaparak eksik/tanımsız adımları tespit eder ve Steps dosyasına ekler."""
        print(f"\n>> 🔍 [Self-Healing] {class_name_hint}Steps için Cucumber adımları denetleniyor...")
        cmd = 'mvn test -Dcucumber.execution.dry-run=true -Dcucumber.plugin="summary"'
        res = subprocess.run(
            cmd,
            cwd=self.java_project_base,
            shell=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="ignore"
        )

        marker = "You can implement missing steps with the snippets below:"
        if marker in res.stdout:
            print(">> ⚠️ [Self-Healing] Tanımlanmamış Cucumber adımları tespit edildi! Yamalanıyor...")
            missing_snippets = res.stdout.split(marker)[1]
            if "```" in missing_snippets:
                missing_snippets = missing_snippets.split("```")[0]
            missing_snippets = re.sub(r"\[INFO\]|\[WARNING\]", "", missing_snippets).strip()

            step_file = os.path.join(self.java_project_root, "StepDefinations", f"{class_name_hint}Steps.java")
            if not os.path.exists(step_file):
                print(f">> Hata: {step_file} bulunamadı!")
                return False

            with open(step_file, "r", encoding="utf-8") as f:
                current_code = f.read()

            prompt = f"""
Sen kıdemli bir Test Otomasyon Mühendisisin.
Aşağıda Cucumber'ın eksik bulduğu adımlar ve mevcut Step sınıfı var.
Mevcut hiçbir metodu bozmadan ve silmeden, sadece bu eksik adımların gövdelerini doldurarak sınıfın sonuna ekle.
TestContext veya Page metotlarını kullanarak metot içlerini anlamlı doldur.

EKSİK ADIMLAR:
{missing_snippets[:3000]}

MEVCUT STEP KODU:
{current_code}

ÇIKTI KURALI:
Sadece tüm sınıfın birleştirilmiş, eksiksiz, çalışan Java kodunu döndür. Markdown veya açıklama yazma.
"""
            healed_code = self.safe_generate(prompt)
            clean_code = re.sub(r"^```[a-zA-Z]*\n|```$", "", healed_code.strip()).strip()

            with open(step_file, "w", encoding="utf-8") as f:
                f.write(clean_code)
            print(">> 🩹 [Self-Healing] Eksik adımlar Steps sınıfına enjekte edildi.")
            return True
        else:
            print(">> ✅ [Self-Healing] Tüm Cucumber adımları sözleşmeye uygun şekilde tanımlı.")
            return False

    def compile_and_heal(self, class_name_hint: str, max_attempts: int = 3) -> bool:
        """Maven derleme kontrolü yapar; hata çıkarsa hedefe yönelik mikro-onarım uygular."""
        for attempt in range(1, max_attempts + 1):
            print(f"\n>> 🔨 [Self-Healing] Derleme Kontrolü (Deneme {attempt}/{max_attempts})...")
            result = subprocess.run(
                "mvn test-compile",
                cwd=self.java_project_base,
                shell=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="ignore"
            )

            if result.returncode == 0:
                print(">> 🎉 [Self-Healing] DERLEME BAŞARILI (BUILD SUCCESS)! Kod mimarisi hatasız.")
                return True

            error_log = result.stdout
            print(f">> ❌ [Self-Healing] Derleme Hatası! Hata logu analiz ediliyor...")
            print(error_log[-1200:])

            heal_prompt = f"""
            Sen uzman bir Java derleme onarıcısısın (Compiler Debugger).
            Maven 'mvn test-compile' sırasında şu hatayı verdi:

            {error_log[-2500:]}

            MİMARİ KURALLAR VE DÜZELTME TALİMATLARI:
            1. BasePage Metot İmzaları:
               - myClick(Locator) -> String parametre kabul etmez, Locator verilmelidir!
               - mySendKeys(Locator, String) -> İlk parametre Locator olmalıdır!
               - myCheckBox(Locator) -> İkinci parametre (boolean vb.) YOKTUR, tek parametre alır!
            2. Step sınıfında çağrılan metot (cannot find symbol: method xyz) Page sınıfında yoksa, Page sınıfına o metodu ekle!
            3. 'incompatible types: String cannot be converted to Locator' hatası varsa, tırnak içindeki seçiciyi 'page.locator("...")' veya ilgili 'get...()' Locator metoduna çevir.
            4. ASLA sınıfları küçültme veya adımları silme.
            5. SADECE hataya sebep olan sınıfı/sınıfları aşağıdaki formatta geri ver:

            ===FILE:Pages/{class_name_hint}Page.java===
            // Düzeltilmiş Page Kodu
            ===END_FILE===

            ===FILE:StepDefinations/{class_name_hint}Steps.java===
            // Düzeltilmiş Step Kodu
            ===END_FILE===
            """
            try:
                healed_bundle = self.safe_generate(heal_prompt)
                pattern = re.compile(r"===FILE:(.*?)===\n(.*?)===END_FILE===", re.DOTALL)
                matches = pattern.findall(healed_bundle)

                for rel_path, code in matches:
                    full_p = os.path.join(self.java_project_root, rel_path.strip())
                    clean_c = re.sub(r"^```[a-zA-Z]*\n|```$", "", code.strip()).strip()
                    os.makedirs(os.path.dirname(full_p), exist_ok=True)
                    with open(full_p, "w", encoding="utf-8") as f:
                        f.write(clean_c)
                    print(f">> 🩹 [Self-Healing] Dosya onarıldı ve güncellendi: {full_p}")
                time.sleep(1)
            except Exception as e:
                print(f">> [Self-Healing İstisna]: {e}")

        return False


def slugify_filename(file_path: str) -> str:
    """Dosya adını temizleyerek feature adına dönüştürür."""
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    clean_name = re.sub(r"[^\w\s-]", "", base_name).strip().lower()
    clean_name = re.sub(r"[-\s]+", "_", clean_name)
    return f"{clean_name}.feature"


def save_to_feature_file(content: str, target_path: str):
    """Üretilen senaryoyu doğrudan .feature dosyasına kaydeder."""
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    with open(target_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"\nDosya başarıyla kaydedildi: {target_path}")



def push_to_github():
    """Değişiklikleri GitHub'a pushlar."""
    print("\n" + "=" * 50)
    print("5. Aşama: Değişiklikler GitHub'a gönderiliyor...")
    print("=" * 50)
    try:
        subprocess.run("git add .", cwd=JAVA_PROJECT_BASE, shell=True, check=True)
        commit_message = f"Otomatik AI pipeline güncellemesi - {time.strftime('%Y-%m-%d %H:%M:%S')}"
        subprocess.run(f'git commit -m "{commit_message}"', cwd=JAVA_PROJECT_BASE, shell=True, check=True)
        subprocess.run("git push origin main", cwd=JAVA_PROJECT_BASE, shell=True, check=True)
        print(">> Değişiklikler başarıyla GitHub'a gönderildi.")
    except subprocess.CalledProcessError as e:
        print(f">> Git işlemi sırasında hata oluştu: {e}")


def get_target_url():
    """URL'i Hooks / Properties dosyasından veya .env'den okur."""
    # 1. Öncelik: Eğer config.properties veya benzeri bir dosyada url tanımlıysa
    if os.path.exists(CONFIG_PROPERTIES_PATH):
        with open(CONFIG_PROPERTIES_PATH, "r", encoding="utf-8") as f:
            for line in f:
                if line.strip().startswith("url") or line.strip().startswith("baseUrl"):
                    return line.split("=")[1].strip()

    # 2. Öncelik: .env dosyasından okuma
    env_url = os.getenv("TARGET_URL")
    if env_url:
        return env_url


# --- 6. ANA ÇALIŞTIRMA BLOĞU ---
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="AI-Agent Test Automation Pipeline")
    parser.add_argument(
        "--mode",
        choices=["design-only", "full-automation"],
        default="design-only",
        help="Çalıştırma modu: 'design-only' (sadece analiz ve Gherkin) veya 'full-automation' (canlı tarama ve kod üretimi)"
    )
    args = parser.parse_args()

    input_folder = "input_docs"
    os.makedirs(input_folder, exist_ok=True)
    os.makedirs(FEATURE_DIR, exist_ok=True)

    # =========================================================================
    # FAZ 1: ERKEN TASARIM (Sitede henüz kod yokken çalışır)
    # =========================================================================
    if args.mode == "design-only":
        print("\n🚀 [FAZ 1] Senaryo ve BDD Tasarım Modu Başlatıldı...")
        input_files = glob.glob(os.path.join(input_folder, "*.docx")) + glob.glob(os.path.join(input_folder, "*.xlsx"))

        if not input_files:
            print(f"❌ Hata: '{input_folder}' klasöründe hiç .docx veya .xlsx dosyası bulunamadı!")
            print(f"Lütfen analiz dosyanızı '{input_folder}' içine atıp tekrar deneyin.")
            sys.exit(1)

        for file_path in input_files:
            print("\n" + "=" * 50)
            print(f">> İşleniyor (Doküman): {file_path}")
            print("=" * 50)

            # 0. Aşama: User Story Çıkarma
            print(">> 0. Aşama: Dokümandan User Story ve İş Kuralları çıkarılıyor...")
            doc_text = extract_content(file_path)
            sample_story = convert_doc_to_user_story(doc_text)

            # 1. Aşama: BDD Gherkin Feature Üretimi & Temizliği
            print(">> 1. Aşama: BDD Gherkin Feature üretiliyor ve normalize ediliyor...")
            raw_feature = generate_bdd_feature(sample_story)
            refined_scenarios = clean_gherkin_syntax(raw_feature)

            # Diske yazma
            feature_filename = slugify_filename(file_path)
            java_feature_path = os.path.join(FEATURE_DIR, feature_filename)
            save_to_feature_file(refined_scenarios, java_feature_path)

        print("\n✅ [FAZ 1 TAMAMLANDI] Senaryolar FeatureFiles klasörüne kalıcı olarak yazıldı.")
        print("⏸️  Geliştirici ekibi sayfayı staging ortamına açana kadar süreç duraklatıldı.\n")
        sys.exit(0)

    # =========================================================================
    # FAZ 2: OTOMASYON & İCRA (Sayfa hazır olduğunda çalışır)
    # =========================================================================
    elif args.mode == "full-automation":
        print("\n⚡ [FAZ 2] Canlı Otomasyon ve Kod Üretim Modu Başlatıldı...")

        feature_files = glob.glob(os.path.join(FEATURE_DIR, "*.feature"))
        if not feature_files:
            print(f"❌ Hata: '{FEATURE_DIR}' dizininde hiç .feature dosyası bulunamadı!")
            print("Lütfen önce 'python scenario_pipeline.py --mode design-only' komutuyla senaryoları üretin.")
            sys.exit(1)

        base_target_url = get_target_url()
        print(f"🔗 Konfigürasyondan alınan temel URL: {base_target_url}")

        for feat_path in feature_files:
            feature_filename = os.path.basename(feat_path)
            print("\n" + "=" * 50)
            print(f">> Otomasyona Alınıyor: {feature_filename}")
            print("=" * 50)

            with open(feat_path, "r", encoding="utf-8") as f:
                refined_scenarios = f.read()

            # 2.5. Aşama: Cucumber CLI ile Resmi Sözleşmeyi Çıkarma
            official_snippets = extract_official_cucumber_snippets(feature_filename)

            # 3. Aşama: PLANNER AGENT (Canlı Keşif ve Eylem Planı)
            planner = PlaywrightPlanner(base_url=base_target_url)
            target_url = planner.resolve_target_url(refined_scenarios)
            live_elements = planner.discover_interactive_elements(target_url)
            execution_plan = planner.create_execution_plan(refined_scenarios, live_elements, safe_generate_content)

            # 4. Aşama: GENERATOR AGENT (Runner, Page, Step Üretimi)
            class_name_hint = feature_filename.replace(".feature", "").title().replace("_", "")
            generator = JavaCodeGenerator(JAVA_PROJECT_ROOT, safe_generate_content)

            # 1. Runner'ı üret
            generator.generate_and_save_runner(class_name_hint)
            # 2. Page sınıfını üret
            page_code = generator.generate_page_class(class_name_hint, execution_plan)
            # 3. Step sınıfını üret
            step_code = generator.generate_step_class(class_name_hint, page_code, official_snippets, refined_scenarios)

            # 5. Aşama: HEALER AGENT (Adım Denetimi ve Derleme Onarımı)
            healer = SelfHealer(JAVA_PROJECT_BASE, JAVA_PROJECT_ROOT, safe_generate_content)
            healer.check_and_heal_missing_cucumber_steps(class_name_hint)
            is_compiled = healer.compile_and_heal(class_name_hint, max_attempts=3)

            # 6. Aşama: Başarılıysa GitHub Push
            if is_compiled:
                print(f"\n>> 🚀 {class_name_hint} İÇİN MİMARİ VE KODLAR ONAYLANDI (BUILD SUCCESS).")
                print(">> Testler GitHub Actions üzerinde çalışmak üzere GitHub'a gönderiliyor...")
                push_to_github()
            else:
                print(f"\n>> 🛑 KRİTİK HATA: {class_name_hint} derlenemedi, GitHub aktarımı durduruldu.")